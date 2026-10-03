#!/usr/bin/env python3
"""Build the four Word reports and the Week 4 PowerPoint presentation.

Run after the weekly analysis scripts and final model training:
    python scripts/build_reports.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUTPUTS = ROOT / "outputs"
PLOTS = OUTPUTS / "plots"
ARTIFACTS = ROOT / "artifacts"

DESCRIPTIONS = {
    "week1": (
        "This Week 1 report documents a complete Python-based data-preparation workflow using "
        "Pandas, NumPy, and scikit-learn. The example is a reproducible synthetic customer-churn "
        "dataset containing 1,800 rows, numeric and categorical input fields, a binary churn "
        "label, and deliberately introduced missing values. The report begins by describing the "
        "dataset structure, data types, label balance, duplicate checks, and missing-value audit. "
        "It then presents exploratory plots and explains why the customer identifier is not used "
        "as a model feature. Numeric gaps are filled with median estimates and categorical "
        "gaps with most-frequent values; missingness indicators are retained for numeric fields. "
        "Categorical variables are one-hot encoded with support for unseen categories, while "
        "numeric inputs are standardized. The churn target is mapped from No and Yes to zero and "
        "one. Mutual information is used to rank the encoded features, and both the complete "
        "transformed table and a top-ten-feature export are included as reproducible artifacts. "
        "The report distinguishes these all-row teaching exports from model evaluation: later "
        "training and cross-validation pipelines learn imputation and scaling only from their "
        "training folds to prevent leakage. Finally, it records the limitations of generated data "
        "and points to the scripts, audit tables, and visualizations needed to repeat the work. "
        "Companion notes introduce Python variables, functions, NumPy arrays, Pandas tables, "
        "model metrics, and leakage concepts in beginner-friendly language."
    ),
    "week2": (
        "This Week 2 report presents a reproducible supervised-learning exercise that covers both "
        "regression and classification. The raw customer-churn sample is split into stratified "
        "training and held-out test portions for the classification task, and every model receives "
        "the same raw inputs through a scikit-learn preprocessing pipeline. Logistic Regression, "
        "a Decision Tree, a Random Forest, and K-Nearest Neighbors predict whether a customer "
        "will churn. Accuracy and balanced accuracy are reported alongside precision, recall, F1, "
        "and ROC-AUC so the comparison reflects both thresholded decisions and probability "
        "ranking. The report explains the observed trade-off: KNN has the largest accuracy in "
        "this run, while Logistic Regression has the strongest ROC-AUC and a higher recall than "
        "KNN; no single score alone establishes the best operational model. Linear Regression is "
        "used for a separate numeric target, monthly charges, rather than being incorrectly applied "
        "to class labels. It is evaluated with MAE, RMSE, and R-squared, and `total_charges` is "
        "excluded because it is derived from the regression target and could leak information. "
        "Tables, a model-comparison chart, code, and saved metric files make the results "
        "auditable. The report also notes the synthetic nature of the data, the fixed random seed, "
        "the held-out test protocol, and the need to use training-only cross-validation for "
        "selection before any real-world deployment."
    ),
    "week3": (
        "This Week 3 report combines unsupervised learning with a careful evaluation of the "
        "supervised churn classifier. K-Means and Ward-linkage Agglomerative Clustering are run "
        "on imputed, scaled, and one-hot-encoded customer features for candidate cluster counts "
        "from two through six. The churn label is deliberately left out of clustering; it is used "
        "only afterward to describe the composition of the selected groups. Silhouette scores "
        "are compared on a reproducible sample of up to 1,000 observations, and a two-component "
        "PCA plot visualizes the selected clustering solution. The report discusses both the "
        "highest-scoring Agglomerative result and its important limitation: the four-cluster "
        "solution is highly imbalanced, so its silhouette value should not be mistaken for proof "
        "of four stable business segments. For model evaluation, a five-fold stratified grid "
        "search tunes Logistic Regression regularization on training data only. The best mean "
        "cross-validation ROC-AUC is approximately 0.802 with a standard deviation of 0.022, and "
        "the selected model is evaluated once on the untouched test split. The report includes "
        "the confusion matrix and accuracy, precision, recall, F1, and ROC-AUC, explaining how "
        "false positives and false negatives differ. It also records the PCA variance, tuning "
        "table, cluster metrics, and limitations of the synthetic sample. Reproducible scripts "
        "and CSV/JSON outputs are supplied so each score and visualization can be regenerated."
    ),
    "week4": (
        "This Week 4 report documents the capstone as a complete local machine-learning "
        "application rather than only a fitted estimator. A Logistic Regression model is trained "
        "using the regularization setting selected by the Week 3 cross-validation exercise, then "
        "serialized with Joblib as a single scikit-learn pipeline. The saved artifact contains "
        "the numeric median imputer, missingness indicators, categorical most-frequent imputer, "
        "standard scaler, one-hot encoder, and classifier, so API callers submit raw customer "
        "fields rather than manually transformed vectors. A FastAPI service exposes a health "
        "route, interactive Swagger documentation, and a batch prediction route with input "
        "validation, missing-field support, unknown-category tolerance, predicted churn labels, "
        "and probabilities. The report describes the request/response contract, local startup "
        "command, test coverage, model metadata, and a worked example. A separate PowerPoint "
        "presentation summarizes the data, model results, clustering, and deployment design. "
        "The final fitted model uses all available rows only after the held-out evaluation "
        "results have been recorded; the test set was not used to choose its tuning parameter. "
        "The implementation is reproducible from the documented commands and includes automated "
        "checks for preprocessing and API behavior. This is a local educational deployment, not "
        "a claim of cloud hosting or production readiness. The dataset and learned relationships "
        "are synthetic; real use would require representative data, privacy controls, authentication, "
        "monitoring, threshold review, and model-version management."
    ),
}


def load_json(name: str) -> dict:
    return json.loads((OUTPUTS / name).read_text(encoding="utf-8"))


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", text))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_table(document: Document, headers: list[str], rows: list[list[str]], font_size: int = 9):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.autofit = True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, text in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17324D")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
    for source_row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(source_row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_picture(document: Document, path: Path, caption: str, width: float = 6.25) -> None:
    if path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        cap = document.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 105, 120)


def base_document(title: str, subtitle: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(41, 52, 63)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Title", 28, "17324D"),
        ("Heading 1", 17, "17324D"),
        ("Heading 2", 12, "237A7A"),
    ]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "YUVA INTERNSHIP  |  PYTHON & MACHINE LEARNING"
    header.style = document.styles["Caption"]
    for run in header.runs:
        run.font.color.rgb = RGBColor(35, 122, 122)
        run.font.bold = True
        run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.text = "Educational capstone • Synthetic customer-churn data • Local reproducible project"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 110, 120)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(25)
    title_run.font.color.rgb = RGBColor(23, 50, 77)
    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(14)
    subtitle_run = p2.add_run(subtitle)
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(35, 122, 122)
    p3 = document.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(20)
    note = p3.add_run("COMPLETED PROJECT REPORT  •  Yuva Internship")
    note.bold = True
    note.font.size = Pt(9)
    note.font.color.rgb = RGBColor(100, 110, 120)
    repository = document.add_paragraph()
    repository.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repository.paragraph_format.space_after = Pt(14)
    repo_run = repository.add_run("Project repository: https://github.com/Dhanam33/Yuva_intern")
    repo_run.font.size = Pt(8)
    repo_run.font.color.rgb = RGBColor(35, 122, 122)
    return document


def add_description(document: Document, key: str) -> None:
    description = DESCRIPTIONS[key]
    count = word_count(description)
    if count < 200:
        raise ValueError(f"{key} description has only {count} words; minimum is 200")
    document.add_heading(f"Submission description ({count} words)", level=1)
    document.add_paragraph(description)


def pct(value: float) -> str:
    return f"{value:.1%}"


def build_week1() -> Path:
    summary = load_json("week1_summary.json")
    missing = pd.read_csv(OUTPUTS / "missing_value_audit.csv")
    features = pd.read_csv(OUTPUTS / "feature_selection.csv").head(10)
    doc = base_document(
        "Week 1 — Data Preprocessing",
        "Python fundamentals • NumPy • Pandas • Exploratory analysis",
    )
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "This report loads, audits, explores, cleans, and transforms a sample dataset for a "
        "machine-learning workflow. It demonstrates numeric and categorical missing-value "
        "handling, encoding, scaling, and feature selection, with reproducible outputs in the repo."
    )
    add_table(
        doc,
        ["Dataset measure", "Result"],
        [
            ["Source", "Synthetic customer-churn sample (seed 2026; not real customer data)"],
            ["Rows", f"{summary['rows']:,}"],
            ["Input features before encoding", str(summary["raw_feature_columns"])],
            ["Raw missing feature cells", f"{summary['raw_missing_feature_cells']:,}"],
            ["Churn labels", f"No {summary['churn_no_count']:,}; Yes {summary['churn_yes_count']:,} ({summary['churn_yes_count'] / summary['rows']:.1%} Yes)"],
            ["Encoded/scaled feature columns", str(summary["encoded_scaled_feature_columns"])],
            ["Missing cells after transformation", str(summary["processed_feature_cells_remaining_missing"])],
        ],
    )
    doc.add_heading("Data loading and quality audit", level=1)
    doc.add_paragraph(
        "The raw CSV contains a customer identifier, 11 model inputs, and a Yes/No target. "
        "The identifier is preserved in the clean export for traceability but excluded from model "
        "inputs. Duplicate rows and duplicate IDs are checked before analysis; none were found "
        "in the seeded sample. The target remains complete so records can be used consistently."
    )
    missing_rows = [
        [row.feature, f"{int(row.missing_count)}", f"{row.missing_percent:.2f}%"]
        for row in missing.itertuples(index=False)
    ]
    add_table(doc, ["Feature", "Missing cells", "Missing %"], missing_rows)
    doc.add_heading("Exploratory data analysis", level=1)
    doc.add_paragraph(
        "The target plot makes the 26.7% positive-class share visible. The contract comparison "
        "shows lower observed churn among longer-term contracts in this generated sample. These "
        "patterns are descriptive, not causal, and should not be generalized to real customers."
    )
    add_picture(doc, PLOTS / "target_distribution.png", "Figure 1. Distribution of the churn target.", width=4.9)
    add_picture(doc, PLOTS / "churn_by_contract.png", "Figure 2. Descriptive churn rate by contract; missing contract values are shown separately.", width=5.7)

    doc.add_heading("Preprocessing decisions", level=1)
    add_bullets(
        doc,
        [
            "Numeric input columns: fill missing values with the median and add indicators for fields that were missing; standardize with StandardScaler.",
            "Categorical input columns: fill missing values with the most frequent value, then apply one-hot encoding with handle_unknown='ignore'.",
            "Target: map No to 0 and Yes to 1. Customer ID is not a predictor; duplicate records/IDs are audited.",
            "Feature selection: rank transformed predictors by mutual information with the churn label and export the top ten for exploratory inspection; this ranking is not used in the model evaluations.",
            "Evaluation safety: the all-row transformed CSV is a Week 1 teaching artifact; later model scripts fit preprocessing only within train/CV folds. Any feature selection used for modeling should also be fitted within those folds.",
        ],
    )
    add_picture(doc, PLOTS / "missing_values.png", "Figure 3. Missingness audit before imputation.", width=5.9)
    doc.add_heading("Top ten features by mutual information", level=2)
    add_table(
        doc,
        ["Rank", "Encoded feature", "Mutual information"],
        [[str(i + 1), row.feature, f"{row.mutual_information:.4f}"] for i, row in enumerate(features.itertuples(index=False))],
    )
    doc.add_heading("Files produced", level=1)
    add_bullets(
        doc,
        [
            "outputs/missing_value_audit.csv — raw counts, percentages, and dtypes.",
            "outputs/preprocessed_customer_churn.csv — complete imputed, scaled, encoded demonstration table.",
            "outputs/selected_feature_dataset.csv and outputs/feature_selection.csv — top-feature export and ranking.",
            "outputs/plots/ — missingness, target balance, and contract/churn visualizations.",
        ],
    )
    doc.add_heading("Limitations", level=1)
    doc.add_paragraph(
        "The dataset is synthetic by design. Median/mode imputation and mutual-information "
        "ranking are baseline teaching choices; a real project would validate assumptions, "
        "consider missingness mechanisms and fairness, and select features using training data only."
    )
    add_description(doc, "week1")
    path = REPORTS / "Week_1_Data_Preprocessing_Report.docx"
    doc.save(path)
    return path


def build_week2() -> Path:
    summary = load_json("week2_summary.json")
    metrics = pd.read_csv(OUTPUTS / "supervised_classification_metrics.csv")
    regression = summary["regression"]
    doc = base_document(
        "Week 2 — Supervised Learning",
        "Regression and classification with scikit-learn",
    )
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "Four classification algorithms are compared on a stratified held-out test split. "
        "Linear Regression is also trained for a continuous target, monthly charges, so both "
        "classification and regression workflows and metrics are demonstrated."
    )
    add_table(
        doc,
        ["Experiment", "Configuration"],
        [
            ["Classification rows", f"{summary['n_train_classification']:,} train / {summary['n_test_classification']:,} test"],
            ["Split", summary["split"]],
            ["Positive class", summary["positive_class"]],
            ["Preprocessing", "Median/mode imputation, numeric scaling, one-hot categories inside each pipeline"],
            ["Regression target", "monthly_charges; total_charges excluded to avoid target leakage"],
        ],
    )
    doc.add_heading("Classification model comparison", level=1)
    doc.add_paragraph(
        "Scores use the same held-out test records. Precision, recall, F1, and ROC-AUC are "
        "especially useful alongside accuracy because churn is the minority class in this sample."
    )
    add_table(
        doc,
        ["Model", "Accuracy", "Balanced acc.", "Precision", "Recall", "F1", "ROC-AUC"],
        [
            [row.model, *(pct(float(getattr(row, column))) for column in ["accuracy", "balanced_accuracy", "precision", "recall", "f1", "roc_auc"])]
            for row in metrics.itertuples(index=False)
        ],
        font_size=8,
    )
    add_picture(doc, PLOTS / "supervised_model_comparison.png", "Figure 1. Precision, recall, F1, and ROC-AUC on the held-out classification test split.", width=6.1)
    logistic = metrics.loc[metrics["model"] == "Logistic Regression"].iloc[0]
    knn = metrics.loc[metrics["model"] == "K-Nearest Neighbors"].iloc[0]
    doc.add_paragraph(
        f"In this run, Logistic Regression has the largest ROC-AUC ({logistic['roc_auc']:.3f}) and "
        f"recall ({logistic['recall']:.3f}) among the four classifiers. KNN has the largest "
        f"accuracy ({knn['accuracy']:.3f}) but lower recall ({knn['recall']:.3f}), illustrating "
        "that predicting the majority class more often can improve accuracy while missing churners. "
        "Threshold and error costs should guide a business decision; the table alone does not establish a winner."
    )
    doc.add_heading("Regression: predict monthly charges", level=1)
    add_table(
        doc,
        ["Model", "Training rows", "Test rows", "MAE", "RMSE", "R-squared"],
        [["Linear Regression", f"{regression['n_train']:,}", f"{regression['n_test']:,}", f"{regression['mae']:.2f}", f"{regression['rmse']:.2f}", f"{regression['r2']:.3f}"]],
    )
    doc.add_paragraph(
        "The regression target has charge units, so MAE and RMSE are interpreted in those units. "
        f"The held-out R-squared is {regression['r2']:.3f}. Total charges are omitted from predictors "
        "because this field is calculated from monthly charges and tenure; using it would make the "
        "exercise unrealistically easy through target leakage."
    )
    doc.add_heading("Method and interpretation", level=1)
    add_bullets(
        doc,
        [
            "The classification test split is stratified and uses a fixed seed (42), preserving approximate class balance.",
            "All four classifier pipelines independently fit imputers, one-hot encoding, and scaling on the training split only.",
            "The reported test metrics are for a fixed comparison. Week 3 performs hyperparameter selection with cross-validation on training data and evaluates the test set only after tuning.",
            "The dataset is synthetic and moderately imbalanced; results are educational and are not expected production performance.",
        ],
    )
    doc.add_heading("Reproducibility", level=1)
    doc.add_paragraph(
        "Run `python -m yuva_ml.week2_supervised`. Results are saved to "
        "outputs/supervised_classification_metrics.csv and outputs/supervised_regression_metrics.json, "
        "with the plotted comparison in outputs/plots/."
    )
    add_description(doc, "week2")
    path = REPORTS / "Week_2_Supervised_Learning_Report.docx"
    doc.save(path)
    return path


def build_week3() -> Path:
    summary = load_json("week3_summary.json")
    evaluation = load_json("tuned_model_evaluation.json")
    clusters = pd.read_csv(OUTPUTS / "clustering_metrics.csv")
    cluster_summary = pd.read_csv(OUTPUTS / "selected_cluster_summary.csv")
    cv = pd.read_csv(OUTPUTS / "cross_validation_results.csv")
    doc = base_document(
        "Week 3 — Unsupervised Learning & Evaluation",
        "Clustering • PCA • cross-validation • model evaluation",
    )
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "This report evaluates K-Means and Agglomerative clustering, uses PCA for visualization, "
        "and tunes a supervised Logistic Regression model with stratified cross-validation. "
        "Clustering excludes the churn target; evaluation metrics use a separate held-out split."
    )
    doc.add_heading("Clustering experiments", level=1)
    doc.add_paragraph(
        "Numeric and categorical features are imputed, scaled, and encoded before clustering. "
        "For each algorithm, candidate values k=2 through k=6 are scored with the silhouette "
        "coefficient on a seeded sample of 1,000 rows. The best score in this run is Agglomerative "
        f"Clustering with k={summary['selected_cluster_solution']['n_clusters']} "
        f"(silhouette {summary['selected_cluster_solution']['silhouette_score']:.3f})."
    )
    add_table(
        doc,
        ["Algorithm", "k", "Silhouette", "Sampled rows"],
        [[row.algorithm, str(int(row.n_clusters)), f"{row.silhouette_score:.3f}", str(int(row.silhouette_sample_size))] for row in clusters.itertuples(index=False)],
    )
    add_picture(doc, PLOTS / "pca_cluster_view.png", "Figure 1. PCA projection colored by the highest-silhouette clustering labels; clustering used the complete encoded feature space.", width=5.7)
    pca = summary["pca"]
    doc.add_paragraph(
        f"The first two principal components explain {pca['two_component_cumulative_explained_variance']:.1%} "
        "of the variance in transformed features. PCA is used only to display a high-dimensional "
        "dataset in two dimensions; clustering itself uses all transformed features."
    )
    doc.add_heading("Cluster composition and caution", level=2)
    add_table(
        doc,
        ["Cluster", "Customers", "Churn count", "Churn rate"],
        [[str(int(row.cluster)), str(int(row.customers)), str(int(row.churn_count)), pct(float(row.churn_rate))] for row in cluster_summary.itertuples(index=False)],
    )
    doc.add_paragraph(
        "The selected four-cluster solution is strongly imbalanced: most observations are in one "
        "cluster while three smaller groups contain only a few dozen records each. The silhouette "
        "score is a geometric separation measure, not proof of useful or stable customer segments. "
        "Cluster sizes, robustness, domain meaning, and alternative distance/scaling choices should "
        "be investigated before acting on any segmentation. Churn rates above are descriptive only."
    )
    doc.add_heading("Cross-validation and hyperparameter tuning", level=1)
    doc.add_paragraph(
        "A five-fold StratifiedKFold grid search evaluates Logistic Regression regularization "
        "values C = 0.1, 1, and 10 using ROC-AUC. Each fold refits preprocessing on its own "
        "training partition. The best parameter is C="
        f"{evaluation['best_params']['classifier__C']:g}, with mean cross-validation ROC-AUC "
        f"{evaluation['best_cross_validation_roc_auc_mean']:.3f} ± "
        f"{evaluation['best_cross_validation_roc_auc_std']:.3f}."
    )
    add_table(
        doc,
        ["C", "Mean CV ROC-AUC", "Std. dev.", "Mean training ROC-AUC"],
        [[f"{row.C:g}", f"{row.mean_cv_roc_auc:.3f}", f"{row.std_cv_roc_auc:.3f}", f"{row.mean_train_roc_auc:.3f}"] for row in cv.itertuples(index=False)],
    )
    doc.add_heading("Held-out test evaluation", level=2)
    test = evaluation["test_metrics"]
    add_table(
        doc,
        ["Accuracy", "Balanced acc.", "Precision", "Recall", "F1", "ROC-AUC"],
        [[pct(test["accuracy"]), pct(test["balanced_accuracy"]), pct(test["precision"]), pct(test["recall"]), pct(test["f1"]), f"{test['roc_auc']:.3f}"]],
    )
    matrix = evaluation["test_confusion_matrix"]
    add_table(
        doc,
        ["Actual / Predicted", "No", "Yes"],
        [["No", str(matrix[0][0]), str(matrix[0][1])], ["Yes", str(matrix[1][0]), str(matrix[1][1])]],
    )
    doc.add_paragraph(
        "The matrix contains true negatives, false positives, false negatives, and true positives "
        "in that order by row. Recall emphasizes finding actual churners; precision reflects how "
        "many flagged customers are actual churners. The test set was not used to tune C. The "
        "metrics characterize this synthetic sample and one fixed random split, not future customers."
    )
    doc.add_heading("Reproducibility", level=1)
    doc.add_paragraph(
        "Run `python -m yuva_ml.week3_unsupervised`. Candidate cluster scores, PCA coordinates, "
        "cluster composition, cross-validation results, the tuned confusion matrix, and JSON "
        "summaries are stored under outputs/."
    )
    add_description(doc, "week3")
    path = REPORTS / "Week_3_Unsupervised_Learning_and_Evaluation_Report.docx"
    doc.save(path)
    return path


def build_week4() -> Path:
    metadata = json.loads((ARTIFACTS / "model_metadata.json").read_text(encoding="utf-8"))
    evaluation = load_json("tuned_model_evaluation.json")
    doc = base_document(
        "Week 4 — AI Deployment & Capstone",
        "Serialized model • FastAPI prediction service • project handoff",
    )
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "The capstone packages preprocessing and the tuned classifier into one Joblib artifact, "
        "then exposes raw-record predictions through a FastAPI application. The API is runnable "
        "locally with interactive documentation and includes basic input validation and tests."
    )
    doc.add_heading("End-to-end design", level=1)
    add_table(
        doc,
        ["Stage", "Implementation"],
        [
            ["Input", "Raw customer fields as JSON; up to 100 records per request; optional/missing values supported"],
            ["Preprocessing", "Median numeric imputation + missing indicators; most-frequent categorical imputation; scaling; one-hot encoding"],
            ["Model", "Logistic Regression; C selected by Week 3 5-fold stratified CV"],
            ["Serialization", "artifacts/churn_model.joblib (full scikit-learn Pipeline)"],
            ["Serving", "FastAPI: GET /health, GET /docs, POST /predict"],
            ["Response", "Predicted label (Yes/No) and churn probability per record"],
        ],
    )
    doc.add_paragraph(
        "The final artifact is refit on all 1,800 available synthetic records after the held-out "
        f"evaluation is recorded. The Week 3 test ROC-AUC was {evaluation['test_metrics']['roc_auc']:.3f}; "
        f"the five-fold mean CV ROC-AUC was {evaluation['best_cross_validation_roc_auc_mean']:.3f} "
        f"± {evaluation['best_cross_validation_roc_auc_std']:.3f}. Those metrics are educational "
        "and do not represent a deployed business outcome."
    )
    doc.add_heading("API contract", level=1)
    add_table(
        doc,
        ["Route", "Purpose", "Result"],
        [
            ["GET /", "Service information", "Name, docs path, prediction route"],
            ["GET /health", "Artifact presence check", "status and model_loaded"],
            ["POST /predict", "Batch inference", "Per-record Yes/No prediction and churn probability"],
            ["GET /docs", "Interactive OpenAPI UI", "Try requests in a browser"],
        ],
    )
    doc.add_heading("Example request", level=2)
    example = (
        '{\n  "instances": [{\n    "customer_id": "demo-001",\n'
        '    "tenure_months": 4, "monthly_charges": 105.0,\n'
        '    "support_tickets": 4, "contract": "Month-to-month",\n'
        '    "internet_service": "Fiber optic",\n'
        '    "payment_method": "Electronic check"\n  }]\n}'
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(example)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(23, 50, 77)
    doc.add_heading("Run and verify", level=1)
    add_bullets(
        doc,
        [
            "Install: python -m pip install -e '.[dev]'",
            "Train/refresh the artifact: python -m yuva_ml.train",
            "Start locally: uvicorn yuva_ml.api:app --app-dir src --host 0.0.0.0 --port 8000",
            "Open http://127.0.0.1:8000/docs or send POST /predict; tests cover health, missing fields, batch predictions, and empty-batch validation.",
            f"Artifact metadata: {metadata['artifact']}; {metadata['training_rows']:,} training records; tuned C={metadata['tuned_regularization_C']:g}.",
        ],
    )
    doc.add_heading("Limitations and next steps", level=1)
    doc.add_paragraph(
        "This is a local educational service, not a cloud-hosted or production-ready application. "
        "The synthetic data cannot validate real churn behavior. Before business use, obtain "
        "representative consented data, assess data quality and subgroup performance, agree on a "
        "cost-sensitive threshold, add authentication and rate limits, monitor drift and service "
        "health, and version the model and dependencies. Joblib artifacts should only be loaded "
        "from trusted sources. The accompanying PowerPoint is reports/Capstone_Presentation.pptx."
    )
    add_description(doc, "week4")
    path = REPORTS / "Week_4_AI_Deployment_Capstone_Report.docx"
    doc.save(path)
    return path


def write_descriptions_markdown() -> Path:
    labels = {
        "week1": "Week 1 — Python for Machine Learning & Data Preprocessing",
        "week2": "Week 2 — Supervised Machine Learning Models",
        "week3": "Week 3 — Unsupervised Learning & Model Evaluation",
        "week4": "Week 4 — AI Project Deployment & Capstone",
    }
    lines = [
        "# Submission descriptions (each 200+ words)",
        "",
        "Copy the relevant description into the internship submission form. Each description is also embedded in its corresponding Word report.",
        "",
    ]
    for key, label in labels.items():
        count = word_count(DESCRIPTIONS[key])
        if count < 200:
            raise ValueError(f"{label} description has only {count} words")
        lines.extend([f"## {label} — {count} words", "", DESCRIPTIONS[key], ""])
    path = REPORTS / "SUBMISSION_DESCRIPTIONS.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def add_ppt_text(slide, x, y, w, h, text, *, size=18, color="17324D", bold=False, align=None):
    shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = PptRGBColor.from_string(color)
    if align is not None:
        paragraph.alignment = align
    return shape


def add_ppt_bullets(slide, x, y, w, h, items, size=18, color="31475E"):
    shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = PptPt(size)
        p.font.color.rgb = PptRGBColor.from_string(color)
        p.level = 0
        p.space_after = PptPt(9)
    return shape


def ppt_slide(prs, title: str, number: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PptRGBColor(248, 250, 252)
    add_ppt_text(slide, 0.65, 0.30, 12.0, 0.55, title, size=26, color="17324D", bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.65), PptInches(0.97), PptInches(12.0), PptInches(0.045))
    line.fill.solid()
    line.fill.fore_color.rgb = PptRGBColor(35, 122, 122)
    line.line.fill.background()
    add_ppt_text(slide, 0.65, 7.12, 11.8, 0.2, "YUVA INTERNSHIP  •  EDUCATIONAL SYNTHETIC-DATA PROJECT", size=8, color="687684")
    add_ppt_text(slide, 12.15, 7.09, 0.5, 0.25, str(number), size=9, color="687684", align=PP_ALIGN.RIGHT)
    return slide


def add_card(slide, x, y, w, h, value, label, accent="237A7A"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PptRGBColor(255, 255, 255)
    shape.line.color.rgb = PptRGBColor.from_string("DCE5EA")
    add_ppt_text(slide, x + 0.12, y + 0.12, w - 0.24, 0.55, value, size=25, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_ppt_text(slide, x + 0.12, y + 0.72, w - 0.24, h - 0.80, label, size=12, color="31475E", align=PP_ALIGN.CENTER)


def build_presentation() -> Path:
    week1 = load_json("week1_summary.json")
    week2 = load_json("week2_summary.json")
    week3 = load_json("week3_summary.json")
    eval_result = load_json("tuned_model_evaluation.json")
    metrics = pd.read_csv(OUTPUTS / "supervised_classification_metrics.csv")
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    # Slide 1: title.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor(23, 50, 77)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(0.22), PptInches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = PptRGBColor(35, 122, 122); accent.line.fill.background()
    add_ppt_text(slide, 0.85, 1.30, 11.2, 0.50, "YUVA INTERNSHIP  •  MACHINE LEARNING", size=14, color="72D6C1", bold=True)
    add_ppt_text(slide, 0.85, 2.00, 11.4, 1.45, "Customer Churn\nPrediction Capstone", size=34, color="FFFFFF", bold=True)
    add_ppt_text(slide, 0.88, 3.85, 10.9, 0.85, "Four-week learning journey: preprocess → compare models → evaluate clusters → deploy an API", size=19, color="D8E3EC")
    add_ppt_text(slide, 0.88, 5.75, 11.2, 0.55, "Built with Python, NumPy, Pandas, scikit-learn, Joblib, and FastAPI", size=14, color="FFFFFF")
    add_ppt_text(slide, 0.88, 6.55, 11.2, 0.3, "Synthetic teaching data only — not a real telecom dataset", size=11, color="AFC3D2")

    # Slide 2: data and preprocessing.
    slide = ppt_slide(prs, "1  |  Data quality and preprocessing", 2)
    add_card(slide, 0.8, 1.35, 3.55, 1.55, f"{week1['rows']:,}", "synthetic customer records")
    add_card(slide, 4.88, 1.35, 3.55, 1.55, f"{week1['raw_missing_feature_cells']:,}", "raw missing feature cells", accent="D17A61")
    add_card(slide, 8.96, 1.35, 3.55, 1.55, str(week1["encoded_scaled_feature_columns"]), "encoded/scaled model inputs")
    add_ppt_bullets(slide, 0.95, 3.45, 5.45, 2.7, [
        "Audited data types, duplicate IDs, missingness, and churn-label balance.",
        "Median imputation + numeric missingness indicators; mode imputation for categories.",
        "StandardScaler for numeric values; one-hot encoding for categorical values.",
        "Mutual information ranks features; top-ten selection is exported.",
    ], size=16)
    if (PLOTS / "missing_values.png").exists():
        slide.shapes.add_picture(str(PLOTS / "missing_values.png"), PptInches(6.65), PptInches(3.15), width=PptInches(5.8))

    # Slide 3: supervised results.
    slide = ppt_slide(prs, "2  |  Supervised learning results", 3)
    rows = len(metrics) + 1
    table_shape = slide.shapes.add_table(rows, 5, PptInches(0.75), PptInches(1.35), PptInches(7.0), PptInches(3.5))
    table = table_shape.table
    headers = ["Model", "Accuracy", "Recall", "F1", "ROC-AUC"]
    for col, heading in enumerate(headers):
        cell = table.cell(0, col); cell.text = heading
    for row_idx, row in enumerate(metrics.itertuples(index=False), start=1):
        values = [row.model, f"{row.accuracy:.3f}", f"{row.recall:.3f}", f"{row.f1:.3f}", f"{row.roc_auc:.3f}"]
        for col, value in enumerate(values):
            table.cell(row_idx, col).text = value
    for row_idx in range(rows):
        for col in range(5):
            cell = table.cell(row_idx, col)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = PptPt(11)
                paragraph.font.name = "Aptos"
                if row_idx == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = PptRGBColor(255, 255, 255)
                    cell.fill.solid(); cell.fill.fore_color.rgb = PptRGBColor(23, 50, 77)
                else:
                    paragraph.font.color.rgb = PptRGBColor(49, 71, 94)
    regression = week2["regression"]
    add_ppt_text(slide, 8.1, 1.45, 4.4, 0.60, "Linear Regression", size=20, color="237A7A", bold=True)
    add_ppt_text(slide, 8.1, 2.05, 4.45, 1.05, f"Monthly charges prediction\nMAE {regression['mae']:.2f}  •  RMSE {regression['rmse']:.2f}\nR² {regression['r2']:.3f}", size=15, color="31475E")
    add_ppt_text(slide, 8.1, 3.45, 4.4, 1.55, "Key readout\nLogistic Regression led ROC-AUC; KNN had the highest accuracy but lower churn recall. Metric choice affects the decision.", size=16, color="31475E")
    add_ppt_text(slide, 0.95, 5.35, 11.4, 0.75, "Classifiers: Logistic Regression • Decision Tree • Random Forest • K-Nearest Neighbors   |   80/20 stratified test split", size=14, color="687684")

    # Slide 4: unsupervised learning and evaluation.
    slide = ppt_slide(prs, "3  |  Clustering and model evaluation", 4)
    solution = week3["selected_cluster_solution"]
    add_card(slide, 0.8, 1.35, 3.45, 1.45, f"{solution['silhouette_score']:.3f}", f"best silhouette • {solution['algorithm']} k={solution['n_clusters']}")
    add_card(slide, 4.72, 1.35, 3.45, 1.45, f"{eval_result['best_cross_validation_roc_auc_mean']:.3f} ± {eval_result['best_cross_validation_roc_auc_std']:.3f}", "5-fold CV ROC-AUC")
    test = eval_result["test_metrics"]
    add_card(slide, 8.64, 1.35, 3.45, 1.45, f"{test['roc_auc']:.3f}", "held-out ROC-AUC")
    if (PLOTS / "pca_cluster_view.png").exists():
        slide.shapes.add_picture(str(PLOTS / "pca_cluster_view.png"), PptInches(0.8), PptInches(3.15), width=PptInches(5.0))
    add_ppt_bullets(slide, 6.9, 3.18, 5.35, 2.85, [
        "K-Means and Agglomerative clustering tested k=2–6; churn was excluded from clustering.",
        f"Tuned Logistic Regression C={eval_result['best_params']['classifier__C']:g} using stratified 5-fold CV.",
        f"Test precision {test['precision']:.3f} • recall {test['recall']:.3f} • F1 {test['f1']:.3f}.",
        "Four-cluster solution is highly imbalanced; silhouette is not proof of natural segments.",
    ], size=15)

    # Slide 5: deployment.
    slide = ppt_slide(prs, "4  |  Deployment architecture", 5)
    steps = [(0.65, "Raw JSON"), (3.05, "FastAPI validation"), (5.45, "Joblib pipeline"), (7.85, "Prediction"), (10.25, "Label + probability")]
    for i, (x, label) in enumerate(steps):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(1.65), PptInches(2.0), PptInches(1.0))
        box.fill.solid(); box.fill.fore_color.rgb = PptRGBColor(255, 255, 255); box.line.color.rgb = PptRGBColor(35, 122, 122)
        add_ppt_text(slide, x + 0.1, 1.82, 1.8, 0.58, label, size=13, color="17324D", bold=True, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_ppt_text(slide, x + 2.01, 1.82, 0.38, 0.55, "→", size=18, color="237A7A", bold=True, align=PP_ALIGN.CENTER)
    add_ppt_text(slide, 0.9, 3.25, 5.8, 0.50, "Local run", size=20, color="237A7A", bold=True)
    add_ppt_text(slide, 0.9, 3.85, 6.0, 1.3, "uvicorn yuva_ml.api:app --app-dir src --host 0.0.0.0 --port 8000\n\nGET /health  •  POST /predict  •  GET /docs", size=13, color="31475E")
    add_ppt_text(slide, 7.15, 3.25, 5.2, 0.5, "Built-in safeguards", size=20, color="237A7A", bold=True)
    add_ppt_bullets(slide, 7.15, 3.8, 5.25, 1.8, [
        "Missing fields imputed; unseen categories tolerated.",
        "Batches limited to 100 rows; empty batches rejected.",
        "Preprocessing travels with the serialized estimator.",
    ], size=15)
    add_ppt_text(slide, 0.9, 5.65, 11.6, 0.75, "Local educational API only — no cloud hosting, authentication, or real-customer validation is claimed.", size=15, color="A2523E", bold=True)

    # Slide 6: conclusion.
    slide = ppt_slide(prs, "5  |  Takeaways and next steps", 6)
    add_ppt_bullets(slide, 1.0, 1.55, 11.1, 3.1, [
        "A reproducible preprocessing pipeline prevents common data-quality and leakage errors.",
        "Multiple metrics expose trade-offs that a single accuracy score can hide.",
        "Unsupervised clusters need balance, stability, and domain validation—not only a score.",
        "A saved end-to-end pipeline makes inference consistent between training and the API.",
    ], size=19)
    add_ppt_text(slide, 1.0, 5.2, 11.2, 0.9, "Next: validate on representative consented data, review costs and fairness, secure the API, and monitor drift.", size=17, color="237A7A", bold=True)

    path = REPORTS / "Capstone_Presentation.pptx"
    prs.save(path)
    return path


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    required = [
        OUTPUTS / "week1_summary.json",
        OUTPUTS / "week2_summary.json",
        OUTPUTS / "week3_summary.json",
        ARTIFACTS / "model_metadata.json",
    ]
    missing = [str(path.relative_to(ROOT)) for path in required if not path.exists()]
    if missing:
        raise SystemExit("Run all analysis and training scripts first; missing: " + ", ".join(missing))
    paths = [build_week1(), build_week2(), build_week3(), build_week4()]
    paths.append(write_descriptions_markdown())
    paths.append(build_presentation())
    print("Generated report deliverables:")
    for path in paths:
        print(f"  {path.relative_to(ROOT)}")
    for key, text in DESCRIPTIONS.items():
        print(f"  {key} description: {word_count(text)} words")


if __name__ == "__main__":
    main()
